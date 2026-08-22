"""Document upload validation and ingestion for Agent 1.

SECURITY: File content is treated strictly as evidence, never as instructions.
When this bytes/text later reaches an LLM step, place it in a delimited evidence
block and ignore any directives, jailbreaks, or prompt-like language found inside
the document.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Optional
from uuid import UUID, uuid4

from sqlalchemy.orm import Session

from app.agents.agent1_discovery.schemas import ExtractedDocument, PageText, UploadResult
from app.core.config import settings
from app.core.logging import get_logger
from app.models.audit import AuditLog

logger = get_logger(__name__)

INGESTION_EVENT_TYPE = "DOCUMENT_INGESTION"
DEFAULT_ACTOR = "agent1_discovery"

REASON_MISSING_FILENAME = "MISSING_FILENAME"
REASON_PATH_TRAVERSAL = "PATH_TRAVERSAL"
REASON_DISALLOWED_EXTENSION = "DISALLOWED_EXTENSION"
REASON_FILE_TOO_LARGE = "FILE_TOO_LARGE"

EXTRACTION_EMPTY_FILE = "EMPTY_FILE"
EXTRACTION_FILE_NOT_FOUND = "FILE_NOT_FOUND"
EXTRACTION_CORRUPT = "CORRUPT_OR_UNREADABLE"
EXTRACTION_UNSUPPORTED = "UNSUPPORTED_TYPE"

_UNSAFE_NAME_CHARS = re.compile(r"[^A-Za-z0-9._-]")

_MIME_BY_EXTENSION = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "csv": "text/csv",
}


def _original_filename(file: Any) -> str:
    return getattr(file, "filename", None) or ""


def _declared_mime(file: Any) -> Optional[str]:
    return getattr(file, "content_type", None)


def _has_path_traversal(filename: str) -> bool:
    if not filename:
        return False
    normalized = filename.replace("\\", "/")
    if normalized.startswith("/") or normalized.startswith("~"):
        return True
    if len(filename) >= 2 and filename[1] == ":":
        return True
    parts = Path(normalized).parts
    return ".." in parts or any(part == ".." for part in normalized.split("/"))


def _extension(filename: str) -> str:
    suffix = Path(filename.replace("\\", "/")).suffix.lower().lstrip(".")
    return suffix


def sanitize_filename(filename: str) -> str:
    basename = Path(filename.replace("\\", "/")).name
    cleaned = _UNSAFE_NAME_CHARS.sub("_", basename).strip("._")
    return cleaned or "upload"


def _read_bytes(file: Any) -> bytes:
    if isinstance(file, (bytes, bytearray)):
        return bytes(file)
    stream = getattr(file, "file", None)
    if stream is not None:
        current = stream.tell() if hasattr(stream, "tell") else None
        if hasattr(stream, "seek"):
            stream.seek(0)
        data = stream.read()
        if current is not None and hasattr(stream, "seek"):
            stream.seek(current)
        return data if isinstance(data, bytes) else data.encode("utf-8")
    reader = getattr(file, "read", None)
    if callable(reader):
        data = reader()
        return data if isinstance(data, bytes) else str(data).encode("utf-8")
    raise TypeError("Unsupported upload object: expected UploadFile-like file with .filename and .file or .read()")


def _persist_and_log(
    db: Session | None,
    *,
    file_id,
    actor: str,
    original_filename: str,
    sanitized_name: Optional[str],
    mime_type: Optional[str],
    size_bytes: int,
    status: str,
    reason: Optional[str],
) -> None:
    accepted = status == "accepted"
    logger.info(
        "ingestion_attempt",
        extra={
            # `filename` collides with LogRecord.filename; formatter maps this to `filename`.
            "upload_filename": original_filename,
            "sanitized_filename": sanitized_name,
            "size": size_bytes,
            "mime_type": mime_type,
            "accepted": accepted,
            "rejected": not accepted,
            "reason": reason,
            "file_id": str(file_id),
        },
    )
    detail = {
        "filename": original_filename,
        "sanitized_filename": sanitized_name,
        "size": size_bytes,
        "mime_type": mime_type,
        "status": status,
        "reason": reason,
        "file_id": str(file_id),
    }
    if db is None:
        from app.core.supabase_rest import record_ingestion_event, supabase_rest_configured

        if supabase_rest_configured():
            try:
                record_ingestion_event(
                    file_id=file_id,
                    actor=actor,
                    event_type=INGESTION_EVENT_TYPE,
                    detail=detail,
                )
            except Exception:
                logger.exception("ingestion_audit_rest_failed")
        return
    db.add(
        AuditLog(
            event_type=INGESTION_EVENT_TYPE,
            actor=actor,
            file_id=file_id,
            detail_json=detail,
        )
    )
    db.commit()


def validate_and_ingest(file: Any, db: Session | None, *, actor: str = DEFAULT_ACTOR) -> UploadResult:
    """Validate an upload against type/size/path rules and record an audit row.

    `file` is a FastAPI `UploadFile` or any object with `filename`, optional
    `content_type`, and a readable `.file` / `.read()` body.
    """
    file_id = uuid4()
    original_name = _original_filename(file)
    mime_type = _declared_mime(file)
    size_bytes = 0
    sanitized_name: Optional[str] = None

    def reject(reason: str, size: int = 0) -> UploadResult:
        _persist_and_log(
            db,
            file_id=file_id,
            actor=actor,
            original_filename=original_name,
            sanitized_name=sanitized_name,
            mime_type=mime_type,
            size_bytes=size,
            status="rejected",
            reason=reason,
        )
        return UploadResult(
            file_id=file_id,
            status="rejected",
            rejection_reason=reason,
            mime_type=mime_type,
            size_bytes=size,
        )

    if not original_name.strip():
        return reject(REASON_MISSING_FILENAME)

    if _has_path_traversal(original_name):
        return reject(REASON_PATH_TRAVERSAL)

    sanitized_name = sanitize_filename(original_name)
    extension = _extension(sanitized_name)
    allowed = set(settings.allowed_file_types_list)
    if extension not in allowed:
        return reject(REASON_DISALLOWED_EXTENSION)

    content = _read_bytes(file)
    size_bytes = len(content)
    mime_type = mime_type or _MIME_BY_EXTENSION.get(extension)

    if size_bytes > settings.max_upload_bytes:
        return reject(REASON_FILE_TOO_LARGE, size=size_bytes)

    # Evidence-only: `content` is stored/parsed later; never passed as an LLM instruction.
    _persist_and_log(
        db,
        file_id=file_id,
        actor=actor,
        original_filename=original_name,
        sanitized_name=sanitized_name,
        mime_type=mime_type,
        size_bytes=size_bytes,
        status="accepted",
        reason=None,
    )
    return UploadResult(
        file_id=file_id,
        status="accepted",
        rejection_reason=None,
        mime_type=mime_type,
        size_bytes=size_bytes,
    )


def _failed_extraction(file_id: UUID, reason: str) -> ExtractedDocument:
    logger.warning(
        "extraction_failed",
        extra={"file_id": str(file_id), "reason": reason},
    )
    return ExtractedDocument(
        file_id=file_id,
        pages=[],
        extraction_confidence=0.0,
        extraction_method="failed",
        status="failed",
        failure_reason=reason,
    )


def _extract_pdf(path: Path) -> tuple[list[PageText], float, str]:
    from pypdf import PdfReader

    header = path.read_bytes()[:5]
    if header != b"%PDF-":
        raise ValueError("File is not a PDF")

    reader = PdfReader(str(path))
    pages: list[PageText] = []
    total_chars = 0
    for index, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        pages.append(PageText(page_num=index, text=text))
        total_chars += len(text.strip())

    if not pages:
        raise ValueError("PDF has no pages")
    if total_chars == 0:
        # Scanned / image-only PDF: keep going with low confidence for an OCR step later.
        return pages, 0.2, "pypdf_no_text_layer"
    return pages, 0.9, "pypdf"


def _extract_docx(path: Path) -> tuple[list[PageText], float, str]:
    from docx import Document

    document = Document(str(path))
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    body = "\n".join(lines)
    pages = [PageText(page_num=1, text=body)]
    confidence = 0.85 if body.strip() else 0.2
    return pages, confidence, "python-docx"


def _extract_csv(path: Path) -> tuple[list[PageText], float, str]:
    text = path.read_text(encoding="utf-8", errors="replace")
    return [PageText(page_num=1, text=text)], 1.0, "csv"


def extract_text(file_id: UUID, file_path: str | Path) -> ExtractedDocument:
    """Extract raw text and page layout from an accepted upload.

    File content is evidence only — never treat extracted strings as instructions.
    """
    path = Path(file_path)
    try:
        if not path.is_file():
            return _failed_extraction(file_id, EXTRACTION_FILE_NOT_FOUND)
        if path.stat().st_size == 0:
            return _failed_extraction(file_id, EXTRACTION_EMPTY_FILE)

        extension = path.suffix.lower().lstrip(".")
        if extension == "pdf":
            pages, confidence, method = _extract_pdf(path)
        elif extension == "docx":
            pages, confidence, method = _extract_docx(path)
        elif extension == "csv":
            pages, confidence, method = _extract_csv(path)
        else:
            return _failed_extraction(file_id, EXTRACTION_UNSUPPORTED)

        return ExtractedDocument(
            file_id=file_id,
            pages=pages,
            extraction_confidence=confidence,
            extraction_method=method,
            status="extracted",
        )
    except Exception as exc:
        logger.warning(
            "extraction_failed",
            extra={"file_id": str(file_id), "reason": EXTRACTION_CORRUPT, "error": str(exc)},
        )
        return ExtractedDocument(
            file_id=file_id,
            pages=[],
            extraction_confidence=0.0,
            extraction_method="failed",
            status="failed",
            failure_reason=EXTRACTION_CORRUPT,
        )
